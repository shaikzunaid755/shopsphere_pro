"""turns a report (from report_service) into csv / excel / pdf bytes"""
import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMATS = {
    "csv": ("text/csv", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "yes" if value else "no"
    return "" if value is None else value


def _header(name):
    return name.replace("_", " ").title()


def to_csv(report):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([_header(c) for c in report["columns"]])
    for row in report["rows"]:
        w.writerow([_cell(row.get(c)) for c in report["columns"]])
    if report["totals"]:
        w.writerow([])
        for k, v in report["totals"].items():
            w.writerow([_header(k), _cell(v)])
    return buf.getvalue().encode("utf-8-sig")  # BOM so excel opens it with the right encoding


def to_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:30]
    ws.append([_header(c) for c in report["columns"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in report["rows"]:
        ws.append([_cell(row.get(c)) for c in report["columns"]])
    if report["totals"]:
        ws.append([])
        for k, v in report["totals"].items():
            ws.append([_header(k), _cell(v)])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report, subtitle=""):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    parts = [Paragraph(f"ShopSphere Pro - {report['title']}", styles["Title"])]
    if subtitle:
        parts.append(Paragraph(subtitle, styles["Normal"]))
    parts.append(Spacer(1, 10))
    data = [[_header(c) for c in report["columns"]]]
    data += [[str(_cell(row.get(c))) for c in report["columns"]] for row in report["rows"]]
    if len(data) == 1:
        data.append(["No data for these filters"] + [""] * (len(report["columns"]) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f4858")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
    ]))
    parts.append(table)
    if report["totals"]:
        parts.append(Spacer(1, 10))
        text = " | ".join(f"<b>{_header(k)}:</b> {_cell(v)}" for k, v in report["totals"].items())
        parts.append(Paragraph(text, styles["Normal"]))
    doc.build(parts)
    return buf.getvalue()


def render(report, fmt, subtitle=""):
    if fmt == "csv":
        return to_csv(report)
    if fmt == "xlsx":
        return to_excel(report)
    return to_pdf(report, subtitle)
